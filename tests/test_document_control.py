from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from tenable_reports.config.profile import (
    DistributionRecipient,
    DocumentPreparationConfig,
    DocumentVersionControlConfig,
)
from tenable_reports.presentation.document_control import append_document_control


def test_document_control_justifies_every_body_cell() -> None:
    document = Document()

    append_document_control(
        document,
        preparation=DocumentPreparationConfig(
            action="Elaboração do documento",
            name="Equipe responsável",
        ),
        version_control=DocumentVersionControlConfig(
            version="1.0",
            affected_sections="Todas",
            change="Elaboração do conteúdo",
            changed_by="Equipe responsável",
        ),
        recipients=(
            DistributionRecipient(
                name="Destinatário padrão",
                organization="Organização responsável",
                email="destinatario@example.invalid",
            ),
        ),
    )

    assert len(document.tables) == 3
    for table in document.tables:
        assert all(
            paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
            for cell in table.rows[0].cells
            for paragraph in cell.paragraphs
        )
        assert all(
            paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
            for cell in table.rows[1].cells
            for paragraph in cell.paragraphs
        )
        assert all(
            paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
            for row in table.rows[2:]
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
