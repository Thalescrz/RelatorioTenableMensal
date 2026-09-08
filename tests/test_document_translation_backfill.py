from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from tenable_reports.application.document_translation_backfill import (
    repair_docx_translation_artifacts,
    translate_docx_narratives,
)


class DocumentTranslationBackfillTests(unittest.TestCase):
    def test_repairs_nested_provider_literals_and_removes_obsolete_failure_note(self) -> None:
        with tempfile.TemporaryDirectory() as directory_name:
            directory = Path(directory_name)
            source = directory / "source.docx"
            target = directory / "repaired.docx"
            document = Document()
            document.add_heading("Descrição:", level=4)
            document.add_paragraph("['Primeira parte traduzida.', 'en']")
            document.add_paragraph(
                "['Segunda parte com [detalhe].', 'en'] "
                "['Terceira parte traduzida.', 'en']"
            )
            document.add_paragraph(
                "['A tradução automática não pôde ser concluída; "
                "o texto original foi preservado.', 'pt-PT']"
            )
            document.add_heading("Mais informações:", level=4)
            document.add_paragraph("https://example.invalid/advisory")
            document.save(source)

            result = repair_docx_translation_artifacts(
                source_path=source,
                output_path=target,
            )
            paragraphs = [paragraph.text for paragraph in Document(target).paragraphs]

        self.assertEqual(result.changed_paragraphs, 3)
        self.assertIn("Primeira parte traduzida.", paragraphs)
        self.assertIn(
            "Segunda parte com [detalhe]. Terceira parte traduzida.",
            paragraphs,
        )
        self.assertNotIn(
            "A tradução automática não pôde ser concluída; "
            "o texto original foi preservado.",
            paragraphs,
        )
        self.assertIn("https://example.invalid/advisory", paragraphs)

    def test_translates_only_labeled_narratives_and_preserves_following_sections(self) -> None:
        with tempfile.TemporaryDirectory() as directory_name:
            directory = Path(directory_name)
            source = directory / "source.docx"
            target = directory / "translated.docx"
            document = Document()
            document.add_heading("5.1. Example vulnerability", level=2)
            document.add_heading("Descrição:", level=4)
            document.add_paragraph("The vulnerability allows remote code execution.")
            document.add_heading("Solução:", level=4)
            document.add_paragraph("Install the vendor security update.")
            document.add_heading("Mais informações:", level=4)
            document.add_paragraph("https://example.invalid/advisory")
            document.add_heading("Host(s) Afetado(s):", level=4)
            document.add_paragraph("host.invalid")
            document.save(source)
            calls: list[tuple[str, str, str]] = []

            def translator(text: str, source_language: str, target_language: str) -> str:
                calls.append((text, source_language, target_language))
                return {
                    "The vulnerability allows remote code execution.": (
                        "A vulnerabilidade permite execução remota de código."
                    ),
                    "Install the vendor security update.": (
                        "Instale a atualização de segurança do fornecedor."
                    ),
                }[text]

            result = translate_docx_narratives(
                source_path=source,
                output_path=target,
                translator=translator,
            )
            paragraphs = [paragraph.text for paragraph in Document(target).paragraphs]

        self.assertEqual(result.target_blocks, 2)
        self.assertEqual(result.changed_paragraphs, 2)
        self.assertEqual(
            calls,
            [
                (
                    "The vulnerability allows remote code execution.",
                    "auto",
                    "pt-BR",
                ),
                ("Install the vendor security update.", "auto", "pt-BR"),
            ],
        )
        self.assertIn("A vulnerabilidade permite execução remota de código.", paragraphs)
        self.assertIn("Instale a atualização de segurança do fornecedor.", paragraphs)
        self.assertIn("https://example.invalid/advisory", paragraphs)
        self.assertIn("host.invalid", paragraphs)

    def test_translation_failure_does_not_publish_partial_document(self) -> None:
        with tempfile.TemporaryDirectory() as directory_name:
            directory = Path(directory_name)
            source = directory / "source.docx"
            target = directory / "translated.docx"
            document = Document()
            document.add_heading("Descrição:", level=4)
            document.add_paragraph("First English sentence. Second English sentence.")
            document.save(source)

            def translator(text: str, _source: str, _target: str) -> str:
                raise RuntimeError("translation unavailable")

            with self.assertRaisesRegex(RuntimeError, "tradução completa"):
                translate_docx_narratives(
                    source_path=source,
                    output_path=target,
                    translator=translator,
                    max_chars=24,
                )

            self.assertFalse(target.exists())

    def test_document_without_narrative_blocks_is_copied_without_repackaging(self) -> None:
        with tempfile.TemporaryDirectory() as directory_name:
            directory = Path(directory_name)
            source = directory / "source.docx"
            target = directory / "translated.docx"
            document = Document()
            document.add_paragraph("Resumo mensal sem detalhamento técnico.")
            document.save(source)

            result = translate_docx_narratives(
                source_path=source,
                output_path=target,
                translator=lambda *_: "não deve ser chamado",
            )

            self.assertEqual(result.target_blocks, 0)
            self.assertEqual(result.changed_paragraphs, 0)
            self.assertEqual(target.read_bytes(), source.read_bytes())


if __name__ == "__main__":
    unittest.main()
