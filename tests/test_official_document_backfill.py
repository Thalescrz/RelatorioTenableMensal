from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from tenable_reports.application.official_document_backfill import (
    OfficialDocumentMetadata,
    compose_official_document,
    finalize_official_document,
    plan_official_document_backfill,
)
from tenable_reports.presentation.full_base_report_docx import (
    _append_official_back_cover,
    _clear_body_after_cover_break,
)


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates/corporate/base-v1.docx"


def _refresh_tool():
    spec = importlib.util.spec_from_file_location(
        "refresh_official_report_documents_for_test",
        ROOT / "tools" / "refresh_official_report_documents.py",
    )
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def _docx(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path


def _manifest(root: Path, documents: list[dict[str, object]]) -> Path:
    path = (
        root
        / "data"
        / "manual"
        / "reports"
        / "client-fixture"
        / "run-fixture"
        / "20260801T000000-20260901T000000"
        / "publication-manifest.json"
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "status": "READY_FOR_CONTROLLED_DISTRIBUTION",
                "client_id": "client-fixture",
                "run_id": "run-fixture",
                "period": {
                    "period_id": "2026-08",
                    "start_at": "2026-08-01T03:00:00Z",
                    "end_at": "2026-09-01T03:00:00Z",
                    "timezone": "America/Fortaleza",
                },
                "documents": documents,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return path


def test_plan_targets_only_documents_cataloged_by_valid_publication_manifests(
    tmp_path: Path,
) -> None:
    report_root = tmp_path / "data" / "manual" / "reports" / "client-fixture"
    paths = {
        kind: _docx(report_root / f"{kind}.docx", f"Documento {kind}")
        for kind in ("base", "custom", "tag", "cloud")
    }
    _docx(tmp_path / "data" / "manual" / "reports" / "orphan.docx", "Órfão")
    _manifest(
        tmp_path,
        [
            {
                "path": str(path),
                "document_kind": kind,
                "document_variant": "expanded" if kind == "cloud" else None,
                "tag_uuid": "tag-fixture" if kind == "tag" else None,
                "tag_category": "Equipe" if kind == "tag" else None,
                "tag_value": "Infraestrutura" if kind == "tag" else None,
            }
            for kind, path in paths.items()
        ],
    )

    plan = plan_official_document_backfill(tmp_path)

    assert plan.manifest_count == 1
    assert plan.document_count == 4
    assert plan.counts_by_kind == {
        "base": 1,
        "cloud": 1,
        "custom": 1,
        "tag": 1,
    }
    assert {item.path for item in plan.manifests[0].documents} == set(paths.values())


def test_plan_rejects_cataloged_document_missing_from_disk(tmp_path: Path) -> None:
    missing = tmp_path / "data" / "manual" / "reports" / "missing.docx"
    _manifest(
        tmp_path,
        [{"path": str(missing), "document_kind": "base"}],
    )

    try:
        plan_official_document_backfill(tmp_path)
    except ValueError as exc:
        assert "não existe no disco" in str(exc)
    else:
        raise AssertionError("Documento catalogado ausente deveria invalidar o plano.")


def _shell_with_body(path: Path, kind: str) -> Path:
    document = Document(TEMPLATE)
    shell = _clear_body_after_cover_break(document)
    if kind == "base":
        document.add_paragraph("CONTROLE DE DOCUMENTO", style="Heading 1")
        document.add_paragraph("OBJETIVO", style="Heading 1")
    elif kind == "custom":
        document.add_paragraph(
            "3.1. Comparativo Mensal de Vulnerabilidades Mitigadas e Não Mitigadas.",
            style="Heading 2",
        )
        document.add_paragraph("Sistemas operacionais e software sem suportes", style="Heading 1")
    elif kind == "tag":
        document.add_paragraph("TAG Equipe - Infraestrutura", style="Heading 1")
        document.add_paragraph("3.2. Principais Ativos Vulneráveis", style="Heading 2")
        document.add_paragraph(
            "VISÃO GERAL DAS PRINCIPAIS VULNERABILIDADES",
            style="Heading 1",
        )
    else:
        document.add_paragraph("1. CONTROLE DE DOCUMENTO")
        document.add_paragraph("3.4. Principais Vulnerabilidades Críticas (TOP 5 CVEs)")
        document.add_paragraph("1. Atualize o pacote para a versão mais recente.")
        document.add_paragraph("4. Conclusão")
    document.add_table(rows=2, cols=2)
    _append_official_back_cover(document, shell)
    document.save(path)
    return path


def _headings(path: Path) -> list[tuple[str, str]]:
    document = Document(path)
    return [
        (paragraph.style.name, paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.style is not None
        and paragraph.style.name.startswith("Heading ")
    ]


def _assert_toc_occupies_its_own_page(document) -> None:
    paragraphs = document.paragraphs
    toc_index = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph.text == "SUMÁRIO"
    )
    first_heading_index = next(
        index for index, paragraph in enumerate(paragraphs[toc_index + 1 :], toc_index + 1)
        if paragraph.style is not None and paragraph.style.name == "Heading 1"
    )
    first_heading = paragraphs[first_heading_index]
    properties = first_heading._p.get_or_add_pPr()
    assert properties.find(qn("w:pageBreakBefore")) is not None
    assert all(
        paragraph._p.find(".//" + qn("w:br")) is None
        for paragraph in paragraphs[toc_index + 1 : first_heading_index]
    )


def test_finalize_official_document_applies_shell_toc_and_numbering_to_every_kind(
    tmp_path: Path,
) -> None:
    expected = {
        "base": ("Heading 1", "1. CONTROLE DE DOCUMENTO"),
        "custom": ("Heading 1", "1. INTELIGÊNCIA E CUSTOMIZAÇÕES TENABLE"),
        "tag": ("Heading 1", "1. TAG Equipe - Infraestrutura"),
        "cloud": ("Heading 2", "3.4. Principais Vulnerabilidades Críticas (TOP 5 CVEs)"),
    }
    for kind in expected:
        path = _shell_with_body(tmp_path / f"{kind}.docx", kind)
        finalize_official_document(
            path,
            OfficialDocumentMetadata(
                document_kind=kind,
                client_name="Cliente Exemplo",
                period_label="AGOSTO/2026",
                period_range="01/08/2026 a 31/08/2026",
                tag_category="Equipe" if kind == "tag" else None,
                tag_value="Infraestrutura" if kind == "tag" else None,
            ),
        )

        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert len(document.sections) == 3
        assert "Cliente Exemplo" in text
        assert "AGOSTO/2026" in text
        assert expected[kind] in _headings(path)
        with zipfile.ZipFile(path) as package:
            document_xml = package.read("word/document.xml").decode("utf-8")
            settings_xml = package.read("word/settings.xml").decode("utf-8")
        assert 'TOC \\o "1-3" \\h \\z' in document_xml
        assert 'TOC \\o "1-4"' not in document_xml
        assert "w:updateFields" in settings_xml
        _assert_toc_occupies_its_own_page(document)
        if kind == "cloud":
            remediation = next(
                paragraph
                for paragraph in document.paragraphs
                if paragraph.text.startswith("1. Atualize o pacote")
            )
            assert not remediation.style.name.startswith("Heading ")


def test_repackaging_validation_accepts_adjacent_tables_merged_by_word(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    source_document = Document()
    source_document.add_paragraph("SUMÁRIO")
    source_document.add_paragraph(
        "Índice estático legado que não pertence ao conteúdo técnico. " * 4
    )
    source_document.add_paragraph("1. CONTROLE DE DOCUMENTO")
    first = source_document.add_table(rows=1, cols=2)
    first.rows[0].cells[0].text = "Campo A"
    first.rows[0].cells[1].text = "Valor A"
    second = source_document.add_table(rows=1, cols=2)
    second.rows[0].cells[0].text = "Campo B"
    second.rows[0].cells[1].text = "Valor B"
    source_document.add_paragraph(
        "SUA MELHOR ALIADA NA JORNADA DA PROTEÇÃO DIGITAL."
    )
    legacy_back_cover = source_document.add_table(rows=1, cols=1)
    legacy_back_cover.cell(0, 0).text = "Rodapé institucional antigo"
    source_document.save(source)

    staged = _shell_with_body(tmp_path / "staged.docx", "cloud")
    staged_document = Document(staged)
    body_table = next(table for table in staged_document.tables if len(table.rows) == 2)
    body_table.rows[0].cells[0].text = "Campo A"
    body_table.rows[0].cells[1].text = "Valor A"
    body_table.rows[1].cells[0].text = "Campo B"
    body_table.rows[1].cells[1].text = "Valor B"
    staged_document.save(staged)

    _refresh_tool()._validate_repackaged(source, staged, document_kind="cloud")


def test_openxml_composition_preserves_body_tables_and_images(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.docx"
    source_document = Document(TEMPLATE)
    shell = _clear_body_after_cover_break(source_document)
    source_document.add_paragraph("SUMÁRIO", style="TOC Heading")
    source_document.add_paragraph("Entrada legada do sumário\t3", style="toc 1")
    source_document.add_page_break()
    source_document.add_paragraph("CONTROLE DE DOCUMENTO", style="Heading 1")
    source_document.add_paragraph("Texto técnico preservado " * 12)
    table = source_document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Campo"
    table.cell(0, 1).text = "Valor"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    image_path = tmp_path / "chart.png"
    Image.new("RGB", (80, 40), "#315efb").save(image_path)
    source_document.add_picture(str(image_path))
    _append_official_back_cover(source_document, shell)
    source_document.save(source)

    output = tmp_path / "composed.docx"
    compose_official_document(
        source=source,
        template=TEMPLATE,
        output=output,
        document_kind="base",
    )

    composed = Document(output)
    assert len(composed.sections) == 3
    assert any("Texto técnico preservado" in paragraph.text for paragraph in composed.paragraphs)
    assert not any("Entrada legada do sumário" in paragraph.text for paragraph in composed.paragraphs)
    assert any(cell.text == "Campo" for item in composed.tables for row in item.rows for cell in row.cells)
    assert _refresh_tool()._body_drawing_count(source, document_kind="base") == 1
    assert _refresh_tool()._body_drawing_count(output, document_kind="base") == 1


def test_technical_text_validation_ignores_legacy_toc_page_numbers(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source-with-legacy-toc.docx"
    source_document = Document(TEMPLATE)
    _clear_body_after_cover_break(source_document)
    source_document.add_paragraph("SUMÁRIO", style="TOC Heading")
    source_document.add_paragraph(
        "Entrada extensa do índice que não pertence ao conteúdo técnico " * 4
        + "\t5",
        style="toc 1",
    )
    technical_text = "Texto técnico real que deve permanecer exatamente igual. " * 4
    source_document.add_paragraph(technical_text)
    source_document.save(source)

    paragraphs = _refresh_tool()._long_body_paragraphs(
        source,
        document_kind="tag",
    )

    assert paragraphs == (" ".join(technical_text.split()),)


def test_repackaging_validation_allows_numbering_change_in_long_heading(
    tmp_path: Path,
) -> None:
    title = "CVE de exemplo com um título técnico deliberadamente extenso " * 3
    source = tmp_path / "source-numbered-heading.docx"
    source_document = Document(TEMPLATE)
    source_shell = _clear_body_after_cover_break(source_document)
    source_document.add_paragraph("SUMÁRIO", style="TOC Heading")
    source_document.add_paragraph(f"5.1. {title}", style="Heading 2")
    _append_official_back_cover(source_document, source_shell)
    source_document.save(source)

    staged = tmp_path / "staged-numbered-heading.docx"
    staged_document = Document(TEMPLATE)
    staged_shell = _clear_body_after_cover_break(staged_document)
    staged_document.add_paragraph("SUMÁRIO", style="TOC Heading")
    staged_document.add_paragraph(f"3.1. {title}", style="Heading 2")
    _append_official_back_cover(staged_document, staged_shell)
    staged_document.save(staged)

    _refresh_tool()._validate_repackaged(source, staged, document_kind="tag")


def test_selected_manifests_skip_the_current_operation_unless_requested(
    tmp_path: Path,
) -> None:
    document_path = _docx(
        tmp_path / "data" / "manual" / "reports" / "current.docx",
        "Documento atual",
    )
    manifest_path = _manifest(
        tmp_path,
        [{"path": str(document_path), "document_kind": "base"}],
    )
    tool = _refresh_tool()
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    payload["document_backfill"] = {
        "operation": tool.OPERATION,
        "document_count": 1,
    }
    manifest_path.write_text(json.dumps(payload), encoding="utf-8")
    manifests = plan_official_document_backfill(tmp_path).manifests

    assert tool._selected_manifests(
        manifests,
        client_id=None,
        limit=None,
        include_completed=False,
    ) == ()
    assert tool._selected_manifests(
        manifests,
        client_id=None,
        limit=None,
        include_completed=True,
    ) == manifests


def test_selected_manifests_target_only_postgresql_main_runs(
    tmp_path: Path,
) -> None:
    document_path = _docx(
        tmp_path / "data" / "manual" / "reports" / "main.docx",
        "Documento MAIN",
    )
    _manifest(
        tmp_path,
        [{"path": str(document_path), "document_kind": "base"}],
    )
    manifests = plan_official_document_backfill(tmp_path).manifests
    tool = _refresh_tool()

    assert tool._selected_manifests(
        manifests,
        client_id=None,
        limit=None,
        include_completed=False,
        main_run_ids=frozenset(),
    ) == ()
    assert tool._selected_manifests(
        manifests,
        client_id=None,
        limit=None,
        include_completed=False,
        main_run_ids=frozenset({"run-fixture"}),
    ) == manifests


def test_word_normalization_repairs_and_updates_only_the_toc_field(
    tmp_path: Path,
) -> None:
    staged = _docx(tmp_path / "staged.docx", "Documento composto")
    script = tmp_path / "compose_official_document.ps1"
    script.write_text("# fixture", encoding="utf-8")
    commands: list[list[str]] = []

    def runner(command, **kwargs):
        commands.append(list(command))
        normalized = Path(command[command.index("-NormalizedOutput") + 1])
        assert normalized.parent == staged.parent
        assert normalized.name.startswith(".word-")
        assert len(normalized.name) < 64
        shutil.copy2(staged, normalized)
        return subprocess.CompletedProcess(command, 0, stdout="UPDATED", stderr="")

    _refresh_tool()._repair_and_update_toc_with_word(
        staged,
        script=script,
        runner=runner,
    )

    assert len(commands) == 1
    assert commands[0][commands[0].index("-Mode") + 1] == "RepairAndUpdateFields"
    assert Document(staged).paragraphs[0].text == "Documento composto"
    assert not list(tmp_path.glob(".word-*.docx"))


def test_word_normalization_repaginates_before_saving_toc_page_numbers() -> None:
    script = (ROOT / "tools" / "compose_official_document.ps1").read_text(
        encoding="utf-8",
    )
    repair_block = script.split('if ($Mode -eq "RepairAndUpdateFields")', 1)[1]
    repair_block = repair_block.split('if (-not $Source -or -not $Template)', 1)[0]

    first_repagination = repair_block.index("$targetDocument.Repaginate()")
    first_update = repair_block.index("$toc.Update()")
    second_repagination = repair_block.index(
        "$targetDocument.Repaginate()",
        first_repagination + 1,
    )
    second_update = repair_block.index("$toc.UpdatePageNumbers()", first_update + 1)

    assert first_repagination < first_update < second_repagination < second_update


def test_word_normalization_reacquires_toc_before_updating_page_numbers() -> None:
    script = (ROOT / "tools" / "compose_official_document.ps1").read_text(
        encoding="utf-8",
    )
    repair_block = script.split('if ($Mode -eq "RepairAndUpdateFields")', 1)[1]
    repair_block = repair_block.split('if (-not $Source -or -not $Template)', 1)[0]

    assert repair_block.count("$targetDocument.TablesOfContents.Item(1)") >= 2
    assert "$toc.UpdatePageNumbers()" in repair_block


def test_staging_uses_short_names_even_when_publication_name_is_long(
    tmp_path: Path,
) -> None:
    source_name = "Relatório de Vulnerabilidades Tenable TAG " + ("Setor-" * 40) + ".docx"

    staged = _refresh_tool()._staged_document_path(
        tmp_path,
        position=3,
        source_name=source_name,
    )

    assert staged == tmp_path / "003.docx"


def test_staging_is_removed_when_document_preparation_fails(
    tmp_path: Path,
    monkeypatch,
) -> None:
    document_path = _docx(
        tmp_path / "data" / "manual" / "reports" / "source.docx",
        "Documento de origem",
    )
    _manifest(
        tmp_path,
        [{"path": str(document_path), "document_kind": "base"}],
    )
    manifest = plan_official_document_backfill(tmp_path).manifests[0]
    tool = _refresh_tool()

    def fail_composition(**kwargs):
        raise RuntimeError("falha controlada")

    monkeypatch.setattr(tool, "compose_official_document", fail_composition)

    with pytest.raises(RuntimeError, match="falha controlada"):
        tool._stage_manifest(
            manifest,
            template=TEMPLATE,
            client_name="Cliente Exemplo",
        )

    assert not list(manifest.path.parent.glob(".official-shell-*"))
