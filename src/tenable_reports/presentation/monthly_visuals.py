from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping, Sequence

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm
from PIL import Image, ImageDraw, ImageFont

from tenable_reports.presentation import base_report_docx as base
from tenable_reports.presentation import full_base_report_docx as faithful


SEVERITY_SERIES = (
    ("critical", "Crítica", "#FF0000"),
    ("high", "Alta", "#ED7D31"),
    ("medium", "Média", "#FFF200"),
    ("low", "Baixa", "#70AD47"),
)


def _font(size: int, bold: bool = False) -> Any:
    candidates = (
        Path(r"C:\Windows\Fonts\calibrib.ttf")
        if bold
        else Path(r"C:\Windows\Fonts\calibri.ttf"),
        Path(r"C:\Windows\Fonts\arialbd.ttf")
        if bold
        else Path(r"C:\Windows\Fonts\arial.ttf"),
    )
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _optional_number(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _format_number(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def _available(row: Mapping[str, Any]) -> bool:
    return str(row.get("availability") or "AVAILABLE") == "AVAILABLE"


def _severity_rows(
    history: Sequence[Mapping[str, Any]],
    *,
    nested_key: str,
    total_key: str,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for item in history:
        row: dict[str, Any] = {
            "label": item.get("label") or item.get("month") or "",
            "availability": item.get("availability") or "AVAILABLE",
            "total": _optional_number(item.get(total_key)) if _available(item) else None,
        }
        nested = item.get(nested_key)
        for key, _, _ in SEVERITY_SERIES:
            row[key] = (
                _optional_number(nested.get(key))
                if _available(item) and isinstance(nested, Mapping)
                else None
            )
        rows.append(row)
    return rows


def render_monthly_table(
    document: Any,
    rows: Sequence[Mapping[str, Any]],
    nested_key: str,
    total_key: str,
) -> Any:
    values: list[tuple[Any, ...]] = []
    for row in rows:
        if not _available(row):
            values.append((row.get("label") or "", "", "", "", "", "Indisponível"))
            continue
        nested = row.get(nested_key)
        nested = nested if isinstance(nested, Mapping) else {}
        values.append(
            (
                row.get("label") or "",
                *(nested.get(key, "") for key, _, _ in SEVERITY_SERIES),
                row.get(total_key, ""),
            )
        )
    return faithful._simple_table(
        document,
        ("Mês", "Crítica", "Alta", "Média", "Baixa", "Total"),
        values,
        widths=(2000, 1250, 1250, 1250, 1250, 1600),
        left_columns=frozenset({0}),
        header_fills=(base.BLUE, base.CRITICAL, base.HIGH, base.MEDIUM, base.LOW, base.BLUE),
    )


def _grouped_monthly_chart(
    path: Path,
    title: str,
    rows: Sequence[Mapping[str, Any]],
    series: Sequence[tuple[str, str, str]],
) -> None:
    width, height = 1500, 900
    image = Image.new("RGB", (width, height), "#292929")
    draw = ImageDraw.Draw(image)
    title_font = _font(38, True)
    label_font = _font(19, True)
    small_font = _font(16)
    draw.text((width // 2, 38), title.upper(), font=title_font, fill="white", anchor="ma")
    left, top, right, bottom = 105, 150, width - 55, height - 205
    values = [
        value
        for row in rows
        for key, _, _ in series
        if (value := _optional_number(row.get(key))) is not None
    ]
    max_value = max(values, default=1) or 1
    for tick in range(6):
        value = round(max_value * tick / 5)
        y = bottom - int((bottom - top) * tick / 5)
        draw.line((left, y, right, y), fill="#595959", width=1)
        draw.text((left - 18, y), _format_number(value), font=small_font, fill="#E7E7E7", anchor="rm")
    group_width = (right - left) / max(len(rows), 1)
    bar_gap = 5
    bar_width = max(12, min(42, int((group_width * 0.78) / max(len(series), 1)) - bar_gap))
    for row_index, row in enumerate(rows):
        group_center = left + group_width * (row_index + 0.5)
        total_width = len(series) * bar_width + (len(series) - 1) * bar_gap
        start_x = int(group_center - total_width / 2)
        for series_index, (key, _, color) in enumerate(series):
            value = _optional_number(row.get(key))
            if value is None:
                continue
            bar_height = int((bottom - top) * value / max_value)
            x1 = start_x + series_index * (bar_width + bar_gap)
            y1 = bottom - bar_height
            draw.rectangle((x1, y1, x1 + bar_width, bottom), fill=color)
            if value:
                draw.text((x1 + bar_width // 2, max(top - 8, y1 - 8)), _format_number(value), font=small_font, fill="white", anchor="ms")
        draw.text((group_center, bottom + 22), str(row.get("label") or ""), font=label_font, fill="white", anchor="ma")
    legend_y = height - 92
    legend_width = sum(42 + draw.textlength(label, font=label_font) + 28 for _, label, _ in series)
    legend_x = int((width - legend_width) / 2)
    for _, label, color in series:
        draw.rectangle((legend_x, legend_y, legend_x + 24, legend_y + 24), fill=color)
        draw.text((legend_x + 34, legend_y + 12), label, font=label_font, fill="white", anchor="lm")
        legend_x += int(42 + draw.textlength(label, font=label_font) + 28)
    image.save(path)


def _monthly_line_chart(
    path: Path,
    title: str,
    history: Sequence[Mapping[str, Any]],
    *,
    value_key: str,
) -> None:
    width, height = 1500, 800
    image = Image.new("RGB", (width, height), "#176783")
    draw = ImageDraw.Draw(image)
    title_font = _font(38, True)
    label_font = _font(18, True)
    small_font = _font(16)
    draw.text((width // 2, 35), title.upper(), font=title_font, fill="white", anchor="ma")
    left, top, right, bottom = 120, 150, width - 70, height - 175
    values = [
        _optional_number(item.get(value_key)) if _available(item) else None
        for item in history
    ]
    max_value = max((value for value in values if value is not None), default=1) or 1
    points: list[tuple[int, int] | None] = []
    for index, value in enumerate(values):
        x = left + int((right - left) * index / max(len(values) - 1, 1))
        if value is None:
            points.append(None)
        else:
            y = bottom - int((bottom - top) * value / max_value)
            points.append((x, y))
            draw.line((x, y, x, bottom), fill="#78AFC1", width=1)
        draw.text((x, bottom + 28), str(history[index].get("label") or ""), font=small_font, fill="white", anchor="ma")
    segment: list[tuple[int, int]] = []
    for point in (*points, None):
        if point is None:
            if len(segment) > 1:
                draw.line(segment, fill="white", width=5, joint="curve")
            segment = []
        else:
            segment.append(point)
    for point, value in zip(points, values):
        if point is None or value is None:
            continue
        x, y = point
        draw.ellipse((x - 6, y - 6, x + 6, y + 6), fill="white")
        draw.text((x, y - 22), _format_number(value), font=label_font, fill="white", anchor="ms")
    image.save(path)


def _add_chart(document: Any, path: Path, alt: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base._add_picture(paragraph, path, width=Cm(16.5), alt_text=alt)


def render_monthly_visual_bundle(
    document: Any,
    rows: Sequence[Mapping[str, Any]],
    output_dir: str | Path,
    scope_label: str,
) -> int:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    year = str(rows[-1].get("label") or "").split("/")[-1] if rows else ""

    faithful._heading(document, "Vulnerabilidades Não Mitigadas", 3)
    render_monthly_table(document, rows, "non_mitigated_by_severity", "non_mitigated")
    non_mitigated_rows = _severity_rows(rows, nested_key="non_mitigated_by_severity", total_key="non_mitigated")
    non_mitigated_comparison = output / "tag-non-mitigated-comparison.png"
    _grouped_monthly_chart(non_mitigated_comparison, f"Comparativo de Vulnerabilidades Não Mitigadas {year}", non_mitigated_rows, (*SEVERITY_SERIES, ("total", "Total Vulnerabilidades", "#B244A5")))
    _add_chart(document, non_mitigated_comparison, f"Comparativo mensal de vulnerabilidades não mitigadas - {scope_label}")
    non_mitigated_volume = output / "tag-non-mitigated-volume.png"
    _monthly_line_chart(non_mitigated_volume, f"Volume de Vuln. Não Mitigadas {year}", rows, value_key="non_mitigated")
    _add_chart(document, non_mitigated_volume, f"Volume mensal de vulnerabilidades não mitigadas - {scope_label}")

    faithful._heading(document, "Vulnerabilidades Mitigadas", 3)
    render_monthly_table(document, rows, "mitigated_by_severity", "mitigated")
    mitigated_rows = _severity_rows(rows, nested_key="mitigated_by_severity", total_key="mitigated")
    mitigated_comparison = output / "tag-mitigated-comparison.png"
    _grouped_monthly_chart(mitigated_comparison, f"Comparativo de Vulnerabilidades Mitigadas {year}", mitigated_rows, (*SEVERITY_SERIES, ("total", "Total Mitigadas", "#B244A5")))
    _add_chart(document, mitigated_comparison, f"Comparativo mensal de vulnerabilidades mitigadas - {scope_label}")
    mitigated_volume = output / "tag-mitigated-volume.png"
    _monthly_line_chart(mitigated_volume, f"Volume de Vuln. Mitigadas {year}", rows, value_key="mitigated")
    _add_chart(document, mitigated_volume, f"Volume mensal de vulnerabilidades mitigadas - {scope_label}")

    faithful._heading(document, "Vulnerabilidades Novas", 3)
    render_monthly_table(document, rows, "new_by_severity", "new")
    evolution_rows = [
        {
            "label": row.get("label") or "",
            "non_mitigated": _optional_number(row.get("non_mitigated")) if _available(row) else None,
            "mitigated": _optional_number(row.get("mitigated")) if _available(row) else None,
            "new": _optional_number(row.get("new")) if _available(row) else None,
        }
        for row in rows
    ]
    evolution = output / "tag-monthly-evolution.png"
    _grouped_monthly_chart(
        evolution,
        f"Evolução Mensal Vulnerabilidades {year}",
        evolution_rows,
        (
            ("non_mitigated", "Total de Vuln. Não Mitigadas", "#FF0000"),
            ("mitigated", "Total Mitigadas", "#00B050"),
            ("new", "Vuln. Novas", "#FFF200"),
        ),
    )
    _add_chart(document, evolution, f"Evolução mensal de vulnerabilidades - {scope_label}")
    return 5


def render_tag_asset_comparison(
    document: Any,
    comparison: Mapping[str, Any],
    *,
    mask_sensitive: bool,
) -> bool:
    periods = [
        item for item in comparison.get("periods") or () if isinstance(item, Mapping)
    ]
    if len(periods) != 2:
        return False
    faithful._heading(document, "Comparativo dos Principais Ativos Vulneráveis", 3)
    for period in periods:
        faithful._paragraph(document, str(period.get("label") or period.get("period_id") or ""), bold=True)
        rows = []
        for index, asset in enumerate(period.get("top_assets") or (), start=1):
            if not isinstance(asset, Mapping):
                continue
            rows.append((
                index,
                "" if mask_sensitive else asset.get("ip_address", ""),
                "" if mask_sensitive else asset.get("asset_name", ""),
                asset.get("critical", ""),
                asset.get("high", ""),
                asset.get("medium", ""),
                asset.get("low", ""),
                asset.get("total", ""),
                asset.get("exploitable", ""),
            ))
        faithful._simple_table(document, ("Nº", "IP Address", "Asset Name", "Crítica", "Alta", "Média", "Baixa", "Total", "Exploitable"), rows, widths=(500, 1350, 1750, 800, 800, 800, 800, 950, 1250), left_columns=frozenset({1, 2}))

    previous_assets = [item for item in periods[0].get("top_assets") or () if isinstance(item, Mapping)]
    current_assets = [item for item in periods[1].get("top_assets") or () if isinstance(item, Mapping)]
    previous_rank = {
        str(item.get("asset_key") or item.get("source_asset_id") or ""): index
        for index, item in enumerate(previous_assets, start=1)
    }
    movement = []
    for current_position, item in enumerate(current_assets, start=1):
        identity = str(item.get("asset_key") or item.get("source_asset_id") or "")
        previous_position = previous_rank.get(identity)
        movement.append((
            "" if mask_sensitive else item.get("ip_address", ""),
            "" if mask_sensitive else item.get("asset_name", ""),
            previous_position if previous_position is not None else "-",
            current_position,
            "Entrada" if previous_position is None else (
                "Subiu" if current_position < previous_position else
                "Desceu" if current_position > previous_position else "Permaneceu"
            ),
        ))
    faithful._simple_table(document, ("IP Address", "Asset Name", "Posição anterior", "Posição atual", "Movimentação"), movement, widths=(1500, 2300, 1500, 1300, 1800), left_columns=frozenset({0, 1, 4}))
    return True

