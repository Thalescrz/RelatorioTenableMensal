from __future__ import annotations

from collections.abc import Callable, Sequence
from typing import Any

from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

from tenable_reports.config.profile import DistributionRecipient
from tenable_reports.presentation import base_report_docx as base


ElementMover = Callable[[Any], Any]


def _move(element: Any, mover: ElementMover | None) -> None:
    if mover is not None:
        mover(element)


def _heading(
    document: DocxDocument,
    text: str,
    *,
    mover: ElementMover | None,
) -> Any:
    paragraph = document.add_paragraph(text, style="Heading 1")
    base._set_paragraph_bottom_border(paragraph, base.BLUE, size=8)
    for run in paragraph.runs:
        base._set_run_font(
            run,
            size=14,
            color=base.NAVY,
            bold=True,
            name="Calibri",
        )
        base._set_language(run)
    _move(paragraph._p, mover)
    return paragraph


def _title_table(
    document: DocxDocument,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    *,
    mover: ElementMover | None,
) -> Any:
    table = document.add_table(rows=2, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    merged = table.cell(0, 0)
    for cell in table.rows[0].cells[1:]:
        merged = merged.merge(cell)
    merged.text = title
    base._set_cell_shading(merged, base.BLUE)
    for run in merged.paragraphs[0].runs:
        base._set_run_font(run, size=8, color=base.WHITE, bold=True)
    merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, header in enumerate(headers):
        cell = table.cell(1, index)
        cell.text = header
        base._set_cell_shading(cell, base.LIGHT_BLUE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            base._set_run_font(run, size=7.2, color=base.NAVY, bold=True)
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            row.cells[index].text = "" if value is None else str(value)
            for run in row.cells[index].paragraphs[0].runs:
                base._set_run_font(run, size=7.2, color=base.NAVY)
        base._prevent_row_split(row)
    _move(table._tbl, mover)
    return table


def _spacer(document: DocxDocument, mover: ElementMover | None) -> None:
    paragraph = document.add_paragraph()
    _move(paragraph._p, mover)


def append_document_control(
    document: DocxDocument,
    *,
    generated_date: str,
    recipients: Sequence[DistributionRecipient],
    mask_sensitive: bool = False,
    mover: ElementMover | None = None,
) -> Any:
    first_heading = _heading(
        document,
        "1. CONTROLE DE DOCUMENTO",
        mover=mover,
    )
    _title_table(
        document,
        "Preparação",
        ("Ação", "Nome", "Data"),
        (("Criação do Documento", "", generated_date),),
        mover=mover,
    )
    _spacer(document, mover)
    _title_table(
        document,
        "Controle de Versionamento",
        ("Versão", "Data da Versão", "Seções Afetadas", "Alteração", "Alterado por"),
        (("1.0", generated_date, "Todas", "Elaboração do conteúdo", ""),),
        mover=mover,
    )
    _spacer(document, mover)
    distribution_rows = tuple(
        (
            "" if mask_sensitive else recipient.name,
            "" if mask_sensitive else recipient.organization,
            "" if mask_sensitive else recipient.email,
        )
        for recipient in recipients
    ) or (("", "", ""), ("", "", ""))
    _title_table(
        document,
        "Lista de Distribuição",
        ("Nome", "Organização", "E-mail"),
        distribution_rows,
        mover=mover,
    )
    return first_heading


__all__ = ["append_document_control"]
