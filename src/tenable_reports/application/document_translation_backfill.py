from __future__ import annotations

import ast
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document

from tenable_reports.application.publishing import validate_docx_package
from tenable_reports.presentation.translation import (
    TextTranslator,
    translate_semantic_text,
)


_TARGET_LABELS = frozenset(
    {
        "descrição:",
        "descricao:",
        "solução:",
        "solucao:",
        "correção ou contramedida recomendada:",
        "correcao ou contramedida recomendada:",
    }
)
_TRANSLATION_FAILURE_NOTE = (
    "a tradução automática não pôde ser concluída; "
    "o texto original foi preservado."
)
_LANGUAGE_CODE = re.compile(r"[a-z]{2,3}(?:-[a-z]{2,4})?", re.IGNORECASE)


@dataclass(frozen=True, slots=True)
class DocumentTranslationResult:
    output_path: Path
    target_blocks: int
    translated_paragraphs: int
    changed_paragraphs: int


def _is_heading(paragraph: Any) -> bool:
    style = getattr(paragraph, "style", None)
    name = str(getattr(style, "name", "") or "").strip().casefold()
    return name.startswith(("heading", "título", "titulo"))


def _replace_text_preserving_format(paragraph: Any, text: str) -> None:
    runs = tuple(paragraph.runs)
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _provider_literal_segments(value: str) -> tuple[str, ...] | None:
    segments: list[str] = []
    index = 0
    while index < len(value):
        while index < len(value) and value[index].isspace():
            index += 1
        if index >= len(value):
            break
        if value[index] != "[":
            return None
        start = index
        depth = 0
        quote = ""
        escaped = False
        while index < len(value):
            character = value[index]
            if quote:
                if escaped:
                    escaped = False
                elif character == "\\":
                    escaped = True
                elif character == quote:
                    quote = ""
            elif character in {"'", '"'}:
                quote = character
            elif character == "[":
                depth += 1
            elif character == "]":
                depth -= 1
                if depth == 0:
                    index += 1
                    break
            index += 1
        if depth != 0 or quote:
            return None
        try:
            decoded = ast.literal_eval(value[start:index])
        except (SyntaxError, ValueError):
            return None
        if (
            not isinstance(decoded, (list, tuple))
            or len(decoded) < 2
            or not isinstance(decoded[0], str)
            or not isinstance(decoded[1], str)
            or _LANGUAGE_CODE.fullmatch(decoded[1].strip()) is None
        ):
            return None
        segments.append(decoded[0].strip())
    return tuple(segments) if segments else None


def _remove_paragraph(paragraph: Any) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def repair_docx_translation_artifacts(
    *,
    source_path: str | Path,
    output_path: str | Path,
) -> DocumentTranslationResult:
    """Remove provider response wrappers left by the legacy Google parser."""

    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    if source == output:
        raise ValueError("O documento reparado precisa usar um caminho de staging.")
    if not source.is_file():
        raise ValueError(f"Documento não encontrado: {source}")

    document = Document(source)
    target_blocks = 0
    changed_paragraphs = 0
    inside_target = False
    for paragraph in tuple(document.paragraphs):
        value = str(paragraph.text or "").strip()
        normalized = value.casefold()
        if normalized in _TARGET_LABELS:
            inside_target = True
            target_blocks += 1
            continue
        if inside_target and _is_heading(paragraph):
            inside_target = False
        if not inside_target or not value:
            continue
        segments = _provider_literal_segments(value)
        if segments is None:
            continue
        repaired = " ".join(
            segment
            for segment in segments
            if segment.casefold() != _TRANSLATION_FAILURE_NOTE
        ).strip()
        if repaired:
            _replace_text_preserving_format(paragraph, repaired)
        else:
            _remove_paragraph(paragraph)
        changed_paragraphs += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    if changed_paragraphs:
        document.save(output)
    else:
        shutil.copy2(source, output)
    validate_docx_package(output)
    return DocumentTranslationResult(
        output_path=output,
        target_blocks=target_blocks,
        translated_paragraphs=0,
        changed_paragraphs=changed_paragraphs,
    )


def translate_docx_narratives(
    *,
    source_path: str | Path,
    output_path: str | Path,
    translator: TextTranslator,
    max_chars: int = 900,
) -> DocumentTranslationResult:
    """Translate only approved narrative blocks in a published report DOCX."""

    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    if source == output:
        raise ValueError("O documento traduzido precisa usar um caminho de staging.")
    if not source.is_file():
        raise ValueError(f"Documento não encontrado: {source}")

    document = Document(source)
    target_blocks = 0
    translated_paragraphs = 0
    changed_paragraphs = 0
    inside_target = False
    for paragraph in document.paragraphs:
        value = str(paragraph.text or "").strip()
        normalized = value.casefold()
        if normalized in _TARGET_LABELS:
            inside_target = True
            target_blocks += 1
            continue
        if inside_target and _is_heading(paragraph):
            inside_target = False
        if not inside_target or not value:
            continue
        translated = translate_semantic_text(
            value,
            translator,
            source_language="auto",
            target_language="pt-BR",
            max_chars=max_chars,
        )
        if translated.had_failures:
            output.unlink(missing_ok=True)
            raise RuntimeError(
                "Não foi possível concluir a tradução completa do documento."
            )
        translated_value = " ".join(translated.chunks).strip()
        translated_paragraphs += 1
        if translated_value and translated_value != value:
            _replace_text_preserving_format(paragraph, translated_value)
            changed_paragraphs += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    if changed_paragraphs:
        document.save(output)
    else:
        shutil.copy2(source, output)
    validate_docx_package(output)
    return DocumentTranslationResult(
        output_path=output,
        target_blocks=target_blocks,
        translated_paragraphs=translated_paragraphs,
        changed_paragraphs=changed_paragraphs,
    )


__all__ = [
    "DocumentTranslationResult",
    "repair_docx_translation_artifacts",
    "translate_docx_narratives",
]
