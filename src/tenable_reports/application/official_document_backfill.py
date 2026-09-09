from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
import json
from pathlib import Path
import re
from typing import Any, Mapping

from docx import Document
from docxcompose.composer import Composer
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from tenable_reports.application.publishing import PublicationDocument
from tenable_reports.presentation import base_report_docx as base
from tenable_reports.presentation import full_base_report_docx as faithful


@dataclass(frozen=True, slots=True)
class OfficialDocumentMetadata:
    document_kind: str
    client_name: str
    period_label: str
    period_range: str
    tag_category: str | None = None
    tag_value: str | None = None


@dataclass(frozen=True, slots=True)
class OfficialBackfillDocument:
    path: Path
    document_kind: str
    document_variant: str | None = None
    tag_uuid: str | None = None
    tag_category: str | None = None
    tag_value: str | None = None

    def publication_document(self) -> PublicationDocument:
        return PublicationDocument(
            self.path,
            self.document_kind,
            document_variant=self.document_variant,
            tag_uuid=self.tag_uuid,
            tag_category=self.tag_category,
            tag_value=self.tag_value,
        )


@dataclass(frozen=True, slots=True)
class OfficialBackfillManifest:
    path: Path
    client_id: str
    run_id: str
    period: Mapping[str, Any]
    documents: tuple[OfficialBackfillDocument, ...]


@dataclass(frozen=True, slots=True)
class OfficialDocumentBackfillPlan:
    manifests: tuple[OfficialBackfillManifest, ...]

    @property
    def manifest_count(self) -> int:
        return len(self.manifests)

    @property
    def document_count(self) -> int:
        return sum(len(item.documents) for item in self.manifests)

    @property
    def counts_by_kind(self) -> dict[str, int]:
        counts = Counter(
            document.document_kind
            for manifest in self.manifests
            for document in manifest.documents
        )
        return dict(sorted(counts.items()))


def _resolved_document_path(manifest: Path, raw: Any) -> Path:
    value = Path(str(raw or ""))
    if not value.is_absolute():
        value = manifest.parent / value
    return value.resolve()


def plan_official_document_backfill(
    project_root: str | Path,
) -> OfficialDocumentBackfillPlan:
    root = Path(project_root).resolve()
    data_root = (root / "data").resolve()
    manifests: list[OfficialBackfillManifest] = []
    seen_documents: set[Path] = set()
    for manifest_path in sorted(data_root.glob("**/reports/**/publication-manifest.json")):
        try:
            payload = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ValueError(f"Manifesto de publicação inválido: {manifest_path}") from exc
        if not isinstance(payload, Mapping):
            raise ValueError(f"Manifesto de publicação inválido: {manifest_path}")
        if str(payload.get("status") or "") != "READY_FOR_CONTROLLED_DISTRIBUTION":
            continue
        client_id = str(payload.get("client_id") or "").strip()
        run_id = str(payload.get("run_id") or "").strip()
        period = payload.get("period")
        raw_documents = payload.get("documents")
        if not client_id or not run_id or not isinstance(period, Mapping):
            raise ValueError(f"Manifesto sem identidade ou período válido: {manifest_path}")
        if not isinstance(raw_documents, list) or not raw_documents:
            raise ValueError(f"Manifesto sem documentos válidos: {manifest_path}")
        documents: list[OfficialBackfillDocument] = []
        for raw_document in raw_documents:
            if not isinstance(raw_document, Mapping):
                raise ValueError(f"Documento inválido no manifesto: {manifest_path}")
            path = _resolved_document_path(manifest_path, raw_document.get("path"))
            try:
                path.relative_to(data_root)
            except ValueError as exc:
                raise ValueError(
                    "Documento catalogado fora da área de dados do projeto."
                ) from exc
            if not path.is_file():
                raise ValueError(f"Documento catalogado não existe no disco: {path}")
            if path in seen_documents:
                raise ValueError(f"Documento catalogado em mais de um manifesto: {path}")
            seen_documents.add(path)
            document = OfficialBackfillDocument(
                path=path,
                document_kind=str(raw_document.get("document_kind") or "").lower(),
                document_variant=(
                    str(raw_document.get("document_variant") or "").lower() or None
                ),
                tag_uuid=str(raw_document.get("tag_uuid") or "").strip() or None,
                tag_category=(
                    str(raw_document.get("tag_category") or "").strip() or None
                ),
                tag_value=str(raw_document.get("tag_value") or "").strip() or None,
            )
            document.publication_document().metadata()
            documents.append(document)
        manifests.append(
            OfficialBackfillManifest(
                path=manifest_path.resolve(),
                client_id=client_id,
                run_id=run_id,
                period=dict(period),
                documents=tuple(documents),
            )
        )
    return OfficialDocumentBackfillPlan(tuple(manifests))


_NUMBER_PREFIX = re.compile(r"^\s*\d+(?:\.\d+)*\.?\s*")


def _without_number(text: str) -> str:
    return _NUMBER_PREFIX.sub("", str(text or "").strip()).strip()


def _set_heading(paragraph: Any, text: str, level: int) -> None:
    paragraph.text = text
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        base._set_run_font(
            run,
            size={1: 14, 2: 11, 3: 10, 4: 9}.get(level, 9),
            color=base.NAVY,
            bold=True,
        )
        base._set_language(run)
    if level == 1:
        base._set_paragraph_bottom_border(paragraph, base.BLUE, size=8)


def _source_body_bounds(document: Any, document_kind: str) -> tuple[int, int]:
    children = list(document._element.body)
    toc_index = next(
        (
            index
            for index, child in enumerate(children)
            if child.tag == qn("w:p")
            and " ".join(Paragraph(child, document).text.split()).casefold()
            == "sumário".casefold()
        ),
        -1,
    )
    if toc_index < 0:
        raise ValueError("Documento de origem sem SUMÁRIO identificável.")

    start_index = -1
    for index, child in enumerate(children[toc_index + 1 :], toc_index + 1):
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document)
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        if document_kind == "cloud":
            if re.match(r"^1\.\s*CONTROLE DE DOCUMENTO$", text, re.IGNORECASE):
                start_index = index
                break
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if not style_name.casefold().startswith("toc"):
            start_index = index
            break
    if start_index < 0:
        raise ValueError("Conteúdo técnico não localizado após o SUMÁRIO.")

    end_index = len(children)
    for index, child in enumerate(children[start_index + 1 :], start_index + 1):
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = " ".join(paragraph.text.split())
            properties = child.find(qn("w:pPr"))
            if "SUA MELHOR ALIADA" in text.upper():
                end_index = index
                break
            if properties is not None and properties.find(qn("w:sectPr")) is not None:
                end_index = index
                break
        elif child.tag == qn("w:sectPr"):
            end_index = index
            break
    if end_index <= start_index:
        raise ValueError("Intervalo técnico inválido no documento de origem.")
    return start_index, end_index


def _retain_technical_body(document: Any, document_kind: str) -> None:
    body = document._element.body
    children = list(body)
    start_index, end_index = _source_body_bounds(document, document_kind)
    retained = set(children[start_index:end_index])
    final_section = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is final_section or child in retained:
            continue
        body.remove(child)


def compose_official_document(
    *,
    source: str | Path,
    template: str | Path,
    output: str | Path,
    document_kind: str,
) -> Path:
    if document_kind not in {"base", "custom", "tag", "cloud"}:
        raise ValueError("Tipo de documento incompatível com o padrão oficial.")
    source_path = Path(source).resolve()
    template_path = Path(template).resolve()
    output_path = Path(output).resolve()
    if not source_path.is_file() or not template_path.is_file():
        raise ValueError("Documento de origem ou template oficial não encontrado.")

    document = Document(template_path)
    shell = faithful._clear_body_after_cover_break(document)
    technical = Document(source_path)
    _retain_technical_body(technical, document_kind)
    composer = Composer(document)
    composer.insert(composer.append_index(), technical)
    faithful._append_official_back_cover(document, shell)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def _normalize_base_headings(paragraphs: list[Any]) -> None:
    titles = (
        ("controle de documento", "1. CONTROLE DE DOCUMENTO"),
        ("objetivo", "2. OBJETIVO"),
        ("sensor nessus, nessus agent e nessus network monitor", "3. SENSOR NESSUS, NESSUS AGENT E NESSUS NETWORK MONITOR"),
        ("visão geral das principais vulnerabilidades", "4. VISÃO GERAL DAS PRINCIPAIS VULNERABILIDADES"),
        ("vulnerabilidades e suas correções e/ou contramedidas recomendadas", "5. VULNERABILIDADES E SUAS CORREÇÕES E/OU CONTRAMEDIDAS RECOMENDADAS"),
        ("sensor was", "6. SENSOR WAS"),
        ("incrementando a segurança e proteção do ambiente", "7. INCREMENTANDO A SEGURANÇA E PROTEÇÃO DO AMBIENTE"),
        ("resumo de vulnerabilidades", "8. RESUMO DE VULNERABILIDADES"),
    )
    for paragraph in paragraphs:
        plain = _without_number(paragraph.text).casefold()
        for key, replacement in titles:
            if plain == key.casefold():
                _set_heading(paragraph, replacement, 1)
                break


def _normalize_tag_headings(paragraphs: list[Any]) -> None:
    section = 0
    detail = 0
    for paragraph in paragraphs:
        plain = _without_number(paragraph.text)
        folded = plain.casefold()
        if folded.startswith("tag "):
            section = 1
            _set_heading(paragraph, f"1. {plain}", 1)
        elif folded == "visão geral das principais vulnerabilidades".casefold():
            section = 2
            _set_heading(paragraph, "2. VISÃO GERAL DAS PRINCIPAIS VULNERABILIDADES", 1)
        elif folded == "vulnerabilidades e suas correções e/ou contramedidas recomendadas".casefold():
            section = 3
            detail = 0
            _set_heading(
                paragraph,
                "3. VULNERABILIDADES E SUAS CORREÇÕES E/OU CONTRAMEDIDAS RECOMENDADAS",
                1,
            )
        elif folded == "comparativo mensal da tag".casefold():
            section = 4
            _set_heading(paragraph, "4. Comparativo Mensal da TAG", 1)
        elif "principais ativos vulneráveis" in folded:
            _set_heading(paragraph, "1.1. Principais Ativos Vulneráveis", 2)
        elif folded == "vulnerabilidades mitigadas":
            _set_heading(paragraph, "2.1. Vulnerabilidades Mitigadas", 2)
        elif folded == "vulnerabilidades não mitigadas":
            _set_heading(paragraph, "2.2. Vulnerabilidades Não Mitigadas", 2)
        elif folded == "vulnerabilidades ressurgidas":
            _set_heading(paragraph, "2.3. Vulnerabilidades Ressurgidas", 2)
        elif (
            section == 3
            and paragraph.style is not None
            and paragraph.style.name == "Heading 2"
        ):
            detail += 1
            _set_heading(paragraph, f"3.{detail}. {plain}", 2)


def _normalize_custom_headings(paragraphs: list[Any]) -> None:
    mappings = (
        ("comparativo mensal de vulnerabilidades mitigadas e não mitigadas", "1.1. Comparativo Mensal de Vulnerabilidades Mitigadas e Não Mitigadas.", 2),
        ("vulnerabilidades “não mitigadas”", "1.1.1. Vulnerabilidades “Não Mitigadas”.", 3),
        ("vulnerabilidades “mitigadas”", "1.1.2. Vulnerabilidades “Mitigadas”.", 3),
        ("integridade da varredura", "1.2. Integridade da varredura", 2),
        ("sistemas operacionais e software sem suportes", "1.5. Sistemas operacionais e softwares sem suporte", 2),
        ("sistemas operacionais e softwares sem suporte", "1.5. Sistemas operacionais e softwares sem suporte", 2),
        ("ativos com sos e softwares sem suportes", "1.5.1. Ativos com SOs e softwares sem suporte.", 3),
        ("ativos com sos e softwares sem suporte", "1.5.1. Ativos com SOs e softwares sem suporte.", 3),
        ("principais softwares e sos sem suportes por vulnerabilidades", "1.5.2. Principais softwares e SOs sem suporte por vulnerabilidades", 3),
        ("principais softwares e sos sem suporte por vulnerabilidades", "1.5.2. Principais softwares e SOs sem suporte por vulnerabilidades", 3),
        ("análise executiva da evolução de vulnerabilidades e criticidade dos ativos", "1.6. Análise Executiva da Evolução de Vulnerabilidades e Criticidade dos Ativos", 2),
        ("evolução mensal de vulnerabilidades", "1.7. Evolução mensal de vulnerabilidades", 2),
        ("tenable cloud security (container images)", "1.8. TENABLE CLOUD SECURITY (CONTAINER IMAGES)", 2),
        ("top 5 imagens de container mais vulneráveis", "1.8.1. Top 5 imagens de container mais vulneráveis", 3),
        ("overview das vulnerabilidades das imagens de container", "1.8.2. Overview das vulnerabilidades das imagens de container", 3),
        ("vulnerabilidades exploráveis por vetor de ataque", "1.9. Vulnerabilidades Exploráveis por Vetor de Ataque", 2),
        ("was vulnerabilidades web – principais aplicações “unsupported”", "1.10. WAS Vulnerabilidades WEB – Principais Aplicações “Unsupported”", 2),
    )
    for paragraph in paragraphs:
        plain = _without_number(paragraph.text).rstrip(".")
        folded = plain.casefold()
        if folded.startswith("comparativo relatório anterior"):
            _set_heading(paragraph, f"1.3. {plain}", 2)
            continue
        for key, replacement, level in mappings:
            if folded == key.casefold():
                _set_heading(paragraph, replacement, level)
                break


def _normalize_cloud_headings(paragraphs: list[Any]) -> None:
    section_titles = (
        "tenable cloud security",
        "introdução",
        "resumo executivo do período",
        "principais hosts vulneráveis",
        "imagens de contêineres mais vulneráveis",
        "overview das vulnerabilidades das imagens de contêiner",
        "principais vulnerabilidades críticas",
        "principais vulnerabilidades com correção disponível",
        "painel de controle",
        "proteção de workloads",
        "status dos sistemas operacionais",
        "componentes e produtos em maior risco",
        "postura de segurança em nuvem",
        "envelhecimento das vulnerabilidades",
        "desempenho de remediação",
        "evolução mensal",
    )
    for paragraph in paragraphs:
        text = str(paragraph.text or "").strip()
        match = re.match(r"^(\d+(?:\.\d+)*)\.?\s+(.+)$", text)
        if match is None:
            continue
        root = match.group(1).split(".", 1)[0]
        if root not in {"1", "2", "3", "4"}:
            continue
        title = match.group(2).strip().casefold()
        already_heading = (
            paragraph.style is not None
            and paragraph.style.name.startswith("Heading ")
        )
        if root == "1" and title != "controle de documento":
            continue
        if root == "2" and title != "objetivo":
            continue
        if root == "4" and not title.startswith("conclus"):
            continue
        if root == "3" and not already_heading and not title.startswith(section_titles):
            continue
        level = min(4, len(match.group(1).split(".")))
        _set_heading(paragraph, text, level)


def _section_paragraphs(document: Any) -> tuple[Any, list[Any], Any]:
    body = document._element.body
    section_breaks = [
        child
        for child in body
        if child.tag == qn("w:p")
        and child.find(qn("w:pPr")) is not None
        and child.find(qn("w:pPr")).find(qn("w:sectPr")) is not None
    ]
    if len(section_breaks) < 2:
        raise ValueError("Documento sem as três seções do padrão oficial.")
    cover_end, body_end = section_breaks[:2]
    children = list(body)
    start_index = children.index(cover_end) + 1
    end_index = children.index(body_end)
    body_nodes = set(children[start_index:end_index])
    paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph._p in body_nodes
    ]
    return cover_end, paragraphs, body_end


def _move_after(anchor: Any, nodes: list[Any]) -> None:
    current = anchor
    for node in nodes:
        current.addnext(node)
        current = node


def finalize_official_document(
    path: str | Path,
    metadata: OfficialDocumentMetadata,
) -> Path:
    destination = Path(path).resolve()
    document = Document(destination)
    if metadata.document_kind not in {"base", "custom", "tag", "cloud"}:
        raise ValueError("Tipo de documento incompatível com o padrão oficial.")
    faithful._configure_styles(document)
    cover_end, paragraphs, _ = _section_paragraphs(document)

    if metadata.document_kind == "base":
        _normalize_base_headings(paragraphs)
    elif metadata.document_kind == "tag":
        _normalize_tag_headings(paragraphs)
    elif metadata.document_kind == "custom":
        _normalize_custom_headings(paragraphs)
    else:
        _normalize_cloud_headings(paragraphs)

    toc_heading = faithful._toc_heading(document)
    faithful._toc_field(document)
    toc_field = document.paragraphs[-1]
    front_nodes = [toc_heading._p, toc_field._p]
    first_heading = next(
        (
            paragraph
            for paragraph in paragraphs
            if paragraph.style is not None and paragraph.style.name == "Heading 1"
        ),
        None,
    )
    if metadata.document_kind == "custom" and not any(
        _without_number(paragraph.text).casefold()
        == "inteligência e customizações tenable".casefold()
        for paragraph in paragraphs
    ):
        title = document.add_paragraph(
            "1. INTELIGÊNCIA E CUSTOMIZAÇÕES TENABLE",
            style="Heading 1",
        )
        _set_heading(title, title.text, 1)
        front_nodes.append(title._p)
        first_heading = title
    if first_heading is None:
        raise ValueError("Documento sem título técnico de nível 1.")
    faithful._page_break_before(first_heading)
    _move_after(cover_end, front_nodes)

    period_label = metadata.period_label
    if metadata.document_kind == "tag":
        tag_label = " - ".join(
            value
            for value in (metadata.tag_category, metadata.tag_value)
            if value
        )
        if tag_label:
            period_label = f"{period_label}\nTAG {tag_label}"
    base._replace_tokens(
        document,
        {
            "{{CLIENT_NAME}}": metadata.client_name,
            "{{PERIOD_LABEL}}": period_label,
            "{{PERIOD_RANGE}}": metadata.period_range,
            "{{TEMPLATE_VERSION}}": faithful.FULL_TEMPLATE_VERSION,
        },
    )
    faithful._sanitize_header_footer(document, metadata.client_name)
    if metadata.document_kind == "custom":
        from tenable_reports.presentation.customizations_report_docx import (
            _cover_title,
        )

        _cover_title(document)
        title = "INTELIGÊNCIA E CUSTOMIZAÇÕES TENABLE"
    elif metadata.document_kind == "cloud":
        from tenable_reports.presentation.cloud_report_docx import (
            _set_cloud_cover_title,
            _set_cloud_header,
        )

        _set_cloud_cover_title(document)
        _set_cloud_header(document, metadata.client_name)
        title = "RELATÓRIO TENABLE CLOUD SECURITY"
    elif metadata.document_kind == "tag":
        title = "RELATÓRIO DE VULNERABILIDADES TENABLE POR TAG"
    else:
        title = faithful.FULL_REPORT_TITLE
    faithful._sanitize_properties(document, title=title)
    base._enable_field_updates(document)
    document.save(destination)
    return destination


__all__ = [
    "OfficialBackfillDocument",
    "OfficialBackfillManifest",
    "OfficialDocumentBackfillPlan",
    "OfficialDocumentMetadata",
    "compose_official_document",
    "finalize_official_document",
    "plan_official_document_backfill",
]
