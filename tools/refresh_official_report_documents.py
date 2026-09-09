from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from tenable_reports.application.official_document_backfill import (  # noqa: E402
    OfficialBackfillDocument,
    OfficialBackfillManifest,
    OfficialDocumentMetadata,
    compose_official_document,
    finalize_official_document,
    plan_official_document_backfill,
)
from tenable_reports.application.publishing import (  # noqa: E402
    PublicationDocumentReplacement,
    refresh_publication_documents_atomically,
    validate_docx_package,
)
from tenable_reports.config.database import DatabaseConfig  # noqa: E402
from tenable_reports.config.environment import load_dotenv_file  # noqa: E402
from tenable_reports.config.profile import load_client_profile  # noqa: E402
from tenable_reports.infrastructure.postgresql import (  # noqa: E402
    PostgresDatabase,
    PostgresOperationsRepository,
)
from tenable_reports.presentation import base_report_docx as base  # noqa: E402


OPERATION = "OFFICIAL_REPORT_SHELL_V3"


def _repair_and_update_toc_with_word(
    staged: Path,
    *,
    script: Path = ROOT / "tools" / "compose_official_document.ps1",
    runner=subprocess.run,
) -> None:
    if not script.is_file():
        raise ValueError("Script de normalização Word não encontrado.")
    normalized = staged.parent / f".word-{uuid.uuid4().hex}.docx"
    command = [
        "powershell.exe",
        "-NoProfile",
        "-ExecutionPolicy",
        "Bypass",
        "-File",
        str(script.resolve()),
        "-Mode",
        "RepairAndUpdateFields",
        "-Output",
        str(staged.resolve()),
        "-NormalizedOutput",
        str(normalized.resolve()),
    ]
    try:
        completed = runner(
            command,
            cwd=str(ROOT),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        if completed.returncode != 0:
            detail = " ".join(
                str(completed.stderr or completed.stdout or "").split()
            )[-500:]
            raise RuntimeError(
                "O Word não conseguiu normalizar e atualizar o sumário do DOCX."
                + (f" Detalhe: {detail}" if detail else "")
            )
        validate_docx_package(normalized)
        normalized.replace(staged)
    finally:
        normalized.unlink(missing_ok=True)


def _table_fingerprint(path: Path) -> tuple[tuple[str, ...], ...]:
    document = Document(path)
    body_children = list(document._element.body)
    toc_index = next(
        (
            index
            for index, child in enumerate(body_children)
            if "SUMÁRIO" in "".join(child.itertext()).upper()
        ),
        -1,
    )
    back_cover_index = next(
        (
            index
            for index, child in enumerate(body_children)
            if index > toc_index
            and "SUA MELHOR ALIADA" in "".join(child.itertext()).upper()
        ),
        len(body_children),
    )
    tables = (
        tuple(
            tuple(cell.text for cell in row.cells)
            for row in Table(child, document).rows
        )
        for index, child in enumerate(body_children)
        if toc_index < index < back_cover_index and child.tag == qn("w:tbl")
    )
    return tuple(
        row
        for table in tables
        if any(cell.strip() for row in table for cell in row)
        for row in table
    )


def _long_body_paragraphs(
    path: Path,
    *,
    document_kind: str,
) -> tuple[str, ...]:
    document = Document(path)
    body_children = list(document._element.body)
    toc_index = next(
        (
            index
            for index, child in enumerate(body_children)
            if "SUMÁRIO" in "".join(child.itertext()).upper()
        ),
        -1,
    )
    start_index = toc_index + 1
    if document_kind == "cloud":
        start_index = next(
            (
                index
                for index, child in enumerate(body_children[start_index:], start_index)
                if child.tag == qn("w:p")
                and re.match(
                    r"^1\.\s*CONTROLE DE DOCUMENTO$",
                    " ".join(Paragraph(child, document).text.split()),
                    flags=re.IGNORECASE,
                )
            ),
            start_index,
        )
    end_index = next(
        (
            index
            for index, child in enumerate(body_children[start_index:], start_index)
            if "SUA MELHOR ALIADA" in "".join(child.itertext()).upper()
        ),
        len(body_children),
    )
    paragraphs = []
    for child in body_children[start_index:end_index]:
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document)
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name.casefold().startswith("toc"):
            continue
        text = _validation_paragraph_text(paragraph)
        if len(text) >= 120 and "SUA MELHOR ALIADA" not in text.upper():
            paragraphs.append(text)
    return tuple(paragraphs)


def _validation_paragraph_text(paragraph: Paragraph) -> str:
    text = " ".join(paragraph.text.split())
    style = paragraph.style
    style_name = style.name if style is not None else ""
    style_id = style.style_id if style is not None else ""
    if (
        style_name.casefold().startswith("heading ")
        or style_id.casefold().startswith("heading")
    ):
        return re.sub(r"^\s*\d+(?:\.\d+)*\.?\s*", "", text).strip()
    return text


def _body_drawing_count(path: Path, *, document_kind: str) -> int:
    document = Document(path)
    if len(document.sections) == 3:
        body = document._element.body
        breaks = [
            child
            for child in body
            if child.tag == qn("w:p")
            and child.find(qn("w:pPr")) is not None
            and child.find(qn("w:pPr")).find(qn("w:sectPr")) is not None
        ]
        if len(breaks) >= 2:
            children = list(body)
            start = children.index(breaks[0]) + 1
            end = children.index(breaks[1])
            return sum(
                len(child.xpath(".//wp:docPr")) for child in children[start:end]
            )
    toc_seen = False
    body_seen = False
    drawings = 0
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        folded = text.casefold()
        if not toc_seen:
            if folded == "sumário".casefold():
                toc_seen = True
            continue
        if not body_seen:
            if document_kind == "cloud":
                body_seen = folded.startswith("1. controle de documento") and len(text) < 100
            else:
                style = paragraph.style.name if paragraph.style is not None else ""
                body_seen = bool(text) and not style.startswith("TOC")
            if not body_seen:
                continue
        if folded.startswith("sua melhor aliada"):
            break
        drawings += len(paragraph._p.xpath(".//wp:docPr"))
    return drawings


def _validate_repackaged(
    source: Path,
    staged: Path,
    *,
    document_kind: str,
) -> None:
    validate_docx_package(staged)
    source_document = Document(source)
    staged_document = Document(staged)
    if len(staged_document.sections) != 3:
        raise ValueError("Documento atualizado não possui as três seções oficiais.")
    source_tables = _table_fingerprint(source)
    staged_tables = _table_fingerprint(staged)
    cursor = iter(staged_tables)
    if not all(any(candidate == expected for candidate in cursor) for expected in source_tables):
        raise ValueError("A republicação alterou o conteúdo de uma ou mais tabelas.")
    if _body_drawing_count(
        source,
        document_kind=document_kind,
    ) != _body_drawing_count(staged, document_kind=document_kind):
        raise ValueError("A republicação alterou gráficos ou imagens do corpo.")
    staged_text = "\n".join(
        _validation_paragraph_text(paragraph)
        for paragraph in staged_document.paragraphs
    )
    missing = [
        paragraph
        for paragraph in _long_body_paragraphs(
            source,
            document_kind=document_kind,
        )
        if paragraph not in staged_text
    ]
    if missing:
        raise ValueError("A republicação não preservou todo o texto técnico do corpo.")


def _metadata(
    manifest: OfficialBackfillManifest,
    document: OfficialBackfillDocument,
    *,
    client_name: str,
) -> OfficialDocumentMetadata:
    period_label, period_range = base._period_labels(manifest.period)
    return OfficialDocumentMetadata(
        document_kind=document.document_kind,
        client_name=client_name,
        period_label=period_label,
        period_range=period_range,
        tag_category=document.tag_category,
        tag_value=document.tag_value,
    )


def _staged_document_path(
    staging: Path,
    *,
    position: int,
    source_name: str,
) -> Path:
    del source_name
    return staging / f"{position:03d}.docx"


def _stage_manifest(
    manifest: OfficialBackfillManifest,
    *,
    template: Path,
    client_name: str,
) -> tuple[Path, tuple[PublicationDocumentReplacement, ...]]:
    staging = manifest.path.parent / f".official-shell-{uuid.uuid4().hex}"
    staging.mkdir(parents=True, exist_ok=False)
    replacements = []
    try:
        for index, document in enumerate(manifest.documents, start=1):
            staged = _staged_document_path(
                staging,
                position=index,
                source_name=document.path.name,
            )
            compose_official_document(
                source=document.path,
                template=template,
                document_kind=document.document_kind,
                output=staged,
            )
            finalize_official_document(
                staged,
                _metadata(manifest, document, client_name=client_name),
            )
            _repair_and_update_toc_with_word(staged)
            _validate_repackaged(
                document.path,
                staged,
                document_kind=document.document_kind,
            )
            replacements.append(
                PublicationDocumentReplacement(
                    staged_path=staged,
                    destination=document.publication_document(),
                )
            )
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    return staging, tuple(replacements)


def _operations(database_env_file: Path) -> PostgresOperationsRepository:
    load_dotenv_file(database_env_file, override=True)
    if not DatabaseConfig.is_configured():
        raise ValueError("A configuração PostgreSQL não está disponível.")
    database = PostgresDatabase(DatabaseConfig.from_environment())
    return PostgresOperationsRepository(database, migrate=False)


def _selected_manifests(
    manifests: Sequence[OfficialBackfillManifest],
    *,
    client_id: str | None,
    limit: int | None,
    include_completed: bool,
    main_run_ids: frozenset[str] | None = None,
) -> tuple[OfficialBackfillManifest, ...]:
    selected = tuple(
        item
        for item in manifests
        if (client_id is None or item.client_id == client_id)
        and (main_run_ids is None or item.run_id in main_run_ids)
        and (include_completed or not _operation_already_applied(item.path))
    )
    return selected if limit is None else selected[:limit]


def _operation_already_applied(manifest_path: Path) -> bool:
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    audit = payload.get("document_backfill")
    return (
        isinstance(audit, dict)
        and str(audit.get("operation") or "") == OPERATION
    )


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Republica somente DOCX de conjuntos MAIN com o padrão oficial, "
            "sem nova coleta."
        )
    )
    parser.add_argument("--project-root", type=Path, default=ROOT)
    parser.add_argument(
        "--template",
        type=Path,
        default=ROOT / "templates" / "corporate" / "base-v1.docx",
    )
    parser.add_argument(
        "--database-env-file",
        type=Path,
        default=ROOT / "credentials" / "database.env",
    )
    parser.add_argument("--client-id")
    parser.add_argument("--limit-manifests", type=int)
    parser.add_argument(
        "--include-completed",
        action="store_true",
        help="Reprocessa também manifestos já marcados com a operação atual.",
    )
    parser.add_argument(
        "--fail-fast",
        action="store_true",
        help="Interrompe no primeiro conjunto que não puder ser republicado.",
    )
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args(argv)

    operations = _operations(args.database_env_file)
    main_run_ids = frozenset(operations.retention_state()["main_run_ids"])
    plan = plan_official_document_backfill(args.project_root)
    manifests = _selected_manifests(
        plan.manifests,
        client_id=args.client_id,
        limit=args.limit_manifests,
        include_completed=args.include_completed,
        main_run_ids=main_run_ids,
    )
    summary = {
        "operation": OPERATION,
        "scope": "postgresql-main",
        "mode": "apply" if args.apply else "dry-run",
        "manifest_count": len(manifests),
        "document_count": sum(len(item.documents) for item in manifests),
        "counts_by_kind": dict(
            sorted(
                __import__("collections").Counter(
                    document.document_kind
                    for manifest in manifests
                    for document in manifest.documents
                ).items()
            )
        ),
    }
    print(json.dumps(summary, ensure_ascii=False))
    if not args.apply:
        return 0
    if not args.template.is_file():
        raise ValueError("Template oficial não encontrado.")
    applied_manifests = 0
    applied_documents = 0
    failures: list[dict[str, object]] = []
    for position, manifest in enumerate(manifests, start=1):
        print(
            json.dumps(
                {
                    "status": "PROCESSING",
                    "manifest": position,
                    "manifest_count": len(manifests),
                    "document_count": len(manifest.documents),
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
        profile_path = args.project_root / "clients" / "managed" / f"{manifest.client_id}.json"
        if not profile_path.is_file():
            raise ValueError("Perfil do cliente catalogado não foi encontrado.")
        profile = load_client_profile(profile_path)
        staging = None
        try:
            staging, replacements = _stage_manifest(
                manifest,
                template=args.template.resolve(),
                client_name=profile.display_name,
            )
            refresh_publication_documents_atomically(
                manifest_path=manifest.path,
                replacements=replacements,
                audit_metadata={
                    "operation": OPERATION,
                    "applied_at": datetime.now(UTC).isoformat(),
                    "document_count": len(replacements),
                },
                commit_callback=lambda path=manifest.path: (
                    operations.record_publication_manifest(path)
                ),
            )
            applied_manifests += 1
            applied_documents += len(replacements)
        except Exception as exc:
            failures.append(
                {
                    "manifest": position,
                    "error_type": type(exc).__name__,
                    "message": str(exc),
                }
            )
            print(
                json.dumps(
                    {
                        "status": "FAILED",
                        "manifest": position,
                        "error_type": type(exc).__name__,
                    },
                    ensure_ascii=False,
                ),
                flush=True,
            )
            if args.fail_fast:
                raise
        finally:
            if staging is not None and staging.is_dir():
                shutil.rmtree(staging, ignore_errors=True)
    print(
        json.dumps(
            {
                "status": "COMPLETE" if not failures else "PARTIAL_FAILURE",
                "manifest_count": applied_manifests,
                "document_count": applied_documents,
                "failed_manifest_count": len(failures),
                "failures": failures,
            },
            ensure_ascii=False,
        )
    )
    return 0 if not failures else 2


if __name__ == "__main__":
    raise SystemExit(main())
